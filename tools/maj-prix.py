#!/usr/bin/env python3
"""Régénère prices.js à partir du tableau de tarifs Excel.

Usage :
    python3 tools/maj-prix.py _sources/tarifs-marine-corail.xlsx

Lit l'onglet « Tarifs » : colonne A = identifiant, colonne L = prix public TTC
en F CFP, colonne M = prix spécial pré-commande TTC (optionnel, F CFP).
Une cellule vide, "Sur demande" ou non reconnue donne null (le catalogue
affiche « Prix sur demande »).
Une cellule au format "À partir de <n>M<d>" (ex. "À partir de 4M5" = 4 500 000)
est acceptée pour les modèles fabriqués/livrés à la commande dont le prix
varie selon la motorisation/les options ; stockée dans prices.js comme
{ "from": <montant> } au lieu d'un nombre, pour que le site affiche
« À partir de X F CFP TTC » plutôt qu'un prix ferme.
Quand la colonne M contient un montant ferme inférieur au prix public de la
même ligne, les deux sont stockés ensemble dans prices.js comme
{ "public": <montant>, "precommande": <montant>, "discount": <pourcentage> } ;
le site n'affiche ce prix double que pour un modèle dont le statut
(inventory.js) est « precommande » (cf. priceHTML() dans app.js), sinon le
prix public normal est affiché. Une colonne M renseignée sans prix public
ferme en face, ou supérieure/égale au prix public, est ignorée et signalée
plutôt que silencieusement acceptée.
Les identifiants absents de vehicles.js sont signalés et ignorés.
Les identifiants qui apparaissent plusieurs fois dans le tableau sont
signalés (le dernier trouvé l'emporte) plutôt que d'être fusionnés en
silence.
Nécessite openpyxl :  pip install openpyxl
"""
import json, sys, re
from pathlib import Path
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent

# "À partir de 4M5" -> 4500000 ; "À partir de 8M" -> 8000000
RE_A_PARTIR_DE = re.compile(r"^\s*à\s*partir\s*de\s*(\d+)\s*[mM]\s*(\d)?\s*$", re.IGNORECASE)

def ids_du_catalogue():
    s = (RACINE / "vehicles.js").read_text(encoding="utf-8")
    data = json.loads(s[s.index("[", s.index("window.")): s.rindex(";")])
    return [v["id"] for v in data]

def parse_prix(valeur):
    """Retourne (montant_ou_None, est_a_partir_de, illisible)."""
    if isinstance(valeur, (int, float)) and valeur > 0:
        return int(valeur), False, False
    if isinstance(valeur, str):
        m = RE_A_PARTIR_DE.match(valeur)
        if m:
            millions, dixiemes = m.group(1), m.group(2)
            montant = int(millions) * 1_000_000 + (int(dixiemes) * 100_000 if dixiemes else 0)
            return montant, True, False
        if valeur.strip().lower() in ("", "sur demande"):
            return None, False, False
        return None, False, True  # texte non reconnu
    return None, False, False

def parse_prix_precommande(valeur):
    """Colonne M : prix spécial pré-commande. Nombre entier ferme uniquement
    (pas de format "à partir de" pour un prix promotionnel). Retourne
    (montant_ou_None, illisible)."""
    if isinstance(valeur, (int, float)) and valeur > 0:
        return int(valeur), False
    if isinstance(valeur, str) and valeur.strip() and valeur.strip().lower() != "sur demande":
        return None, True  # texte non reconnu
    return None, False

def main(xlsx):
    connus = ids_du_catalogue()
    ws = load_workbook(xlsx, data_only=True)["Tarifs"]
    prix, inconnus, doublons, a_partir_de, illisibles = {}, [], {}, [], []
    doubles, precommande_ignores = [], []
    vus = {}
    for n_ligne, ligne in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        ident, valeur = ligne[0], ligne[11]
        valeur_precommande = ligne[12] if len(ligne) > 12 else None
        if not ident or not isinstance(ident, str):
            continue
        ident = ident.strip()
        if not re.match(r"^[A-Za-z0-9-]+$", ident):
            continue
        if ident not in connus:
            inconnus.append((n_ligne, ident)); continue

        if ident in vus:
            doublons.setdefault(ident, [vus[ident]]).append((n_ligne, valeur))
        vus[ident] = (n_ligne, valeur)

        montant, est_a_partir_de, illisible = parse_prix(valeur)
        if illisible:
            illisibles.append((n_ligne, ident, valeur))

        montant_precommande, illisible_precommande = parse_prix_precommande(valeur_precommande)
        if illisible_precommande:
            precommande_ignores.append((n_ligne, ident, valeur_precommande, "valeur non reconnue"))

        if montant_precommande is not None:
            if est_a_partir_de or montant is None:
                precommande_ignores.append((n_ligne, ident, montant_precommande,
                                             "pas de prix public ferme en face pour calculer la remise"))
            elif montant_precommande >= montant:
                precommande_ignores.append((n_ligne, ident, montant_precommande,
                                             f"supérieur ou égal au prix public ({montant})"))
            else:
                remise = round((montant - montant_precommande) / montant * 100)
                prix[ident] = {"public": montant, "precommande": montant_precommande, "discount": remise}
                doubles.append((ident, montant, montant_precommande, remise))
                continue

        if est_a_partir_de:
            a_partir_de.append((ident, montant))
            prix[ident] = {"from": montant}
        else:
            prix[ident] = montant

    manquants = [i for i in connus if i not in prix]
    for i in manquants:
        prix[i] = None
    prix = {i: prix[i] for i in connus}          # ordre du catalogue

    (RACINE / "prices.js").write_text(
        "// Prix Marine Corail en F CFP TTC.\n"
        "// - nombre entier : prix ferme.\n"
        "// - null : prix non renseigné (affiche « Prix sur demande »).\n"
        "// - { \"from\": <nombre> } : prix indicatif (affiche « À partir de X F CFP TTC »),\n"
        "//   pour les modèles fabriqués/livrés à la commande dont le tarif dépend de la\n"
        "//   motorisation/des options.\n"
        "// - { \"public\": <nombre>, \"precommande\": <nombre>, \"discount\": <pourcentage> } :\n"
        "//   prix public + prix spécial pré-commande (colonne M du tableur). N'est affiché\n"
        "//   en double sur le site que si le statut (inventory.js) du modèle est\n"
        "//   \"precommande\" ; sinon le prix public seul est utilisé (cf. priceHTML() dans app.js).\n"
        "// Fichier généré par tools/maj-prix.py à partir de _sources/tarifs-marine-corail.xlsx.\n\n"
        "window.MC_PRICES = " + json.dumps(prix, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8")

    renseignes = sum(1 for v in prix.values() if v is not None)
    print(f"prices.js mis à jour : {renseignes}/{len(prix)} modèles tarifés")

    if a_partir_de:
        print(f"Prix « à partir de » (affichés comme tels sur le site, {len(a_partir_de)}) :")
        for ident, montant in a_partir_de:
            print(f"  - {ident} : à partir de {montant}")

    if doubles:
        print(f"Prix doubles public/pré-commande calculés ({len(doubles)}) :")
        for ident, pub, pre, remise in doubles:
            print(f"  - {ident} : {pub} F -> {pre} F (-{remise}%)")

    if precommande_ignores:
        print("Colonne M (prix pré-commande) ignorée sur certaines lignes :")
        for n_ligne, ident, valeur, raison in precommande_ignores:
            print(f"  - ligne {n_ligne} ({ident}) : {valeur!r} — {raison}")

    if doublons:
        print("Identifiants présents plusieurs fois dans le tableau (dernière valeur conservée) :")
        for ident, occ in doublons.items():
            print(f"  - {ident} :", ", ".join(f"ligne {n} = {v!r}" for n, v in occ))

    if illisibles:
        print("Valeurs de prix non reconnues (traitées comme « Prix sur demande ») :")
        for n_ligne, ident, valeur in illisibles:
            print(f"  - ligne {n_ligne} ({ident}) : {valeur!r}")

    if inconnus:
        print("Identifiants du tableau absents du catalogue (ignorés) :")
        for n_ligne, i in inconnus:
            print(f"  - ligne {n_ligne} : {i}")

    if manquants:
        print("Modèles du catalogue absents du tableau (prix vidé) :")
        for i in manquants: print("  -", i)

if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
