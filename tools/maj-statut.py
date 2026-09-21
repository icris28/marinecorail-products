#!/usr/bin/env python3
"""Met à jour le statut de stock dans inventory.js à partir du tableau tarifs Excel.

Usage :
    python3 tools/maj-statut.py _sources/tarifs-marine-corail.xlsx

Lit l'onglet « Tarifs » : colonne A = identifiant, colonne K = Statut.
Ne touche que le champ "status" de chaque entrée déjà présente dans
inventory.js (quantity/color et le reste du fichier ne sont pas modifiés).
Correspondance des statuts (à ajuster ici si besoin) :
    "En stock"       -> "stock"
    "Pré-commande"   -> "precommande"
    "Sur commande"   -> "commande"
    "Non disponible" -> "rupture"
Un identifiant du tableau absent d'inventory.js, ou un statut texte non
reconnu, est signalé et ignoré (pas de modification en silence).
Nécessite openpyxl :  pip install openpyxl
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent

STATUT_MAP = {
    "en stock": "stock",
    "pré-commande": "precommande",
    "precommande": "precommande",
    "sur commande": "commande",
    "non disponible": "rupture",
}

def ids_du_catalogue():
    s = (RACINE / "vehicles.js").read_text(encoding="utf-8")
    data = json.loads(s[s.index("[", s.index("window.")): s.rindex(";")])
    return [v["id"] for v in data]

def main(xlsx):
    connus = set(ids_du_catalogue())
    ws = load_workbook(xlsx, data_only=True)["Tarifs"]

    statuts, inconnus, doublons, non_reconnus = {}, [], {}, []
    vus = {}
    for n_ligne, ligne in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        ident, statut_txt = ligne[0], ligne[10]
        if not ident or not isinstance(ident, str):
            continue
        ident = ident.strip()
        if not re.match(r"^[A-Za-z0-9-]+$", ident):
            continue
        if ident not in connus:
            inconnus.append((n_ligne, ident)); continue
        if not statut_txt or not isinstance(statut_txt, str):
            continue

        if ident in vus:
            doublons.setdefault(ident, [vus[ident]]).append((n_ligne, statut_txt))
        vus[ident] = (n_ligne, statut_txt)

        cle = statut_txt.strip().lower()
        if cle not in STATUT_MAP:
            non_reconnus.append((n_ligne, ident, statut_txt))
            continue
        statuts[ident] = STATUT_MAP[cle]

    path = RACINE / "inventory.js"
    src = path.read_text(encoding="utf-8")

    changements, absents_du_fichier = [], []
    for ident, nouveau in statuts.items():
        m = re.search(r'"' + re.escape(ident) + r'":\s*\{([^{}]*)\}', src)
        if not m:
            absents_du_fichier.append(ident)
            continue
        bloc = m.group(1)
        m2 = re.search(r'"status":\s*"([^"]*)"', bloc)
        if not m2:
            absents_du_fichier.append(ident)
            continue
        ancien = m2.group(1)
        if ancien == nouveau:
            continue
        nouveau_bloc = bloc[:m2.start()] + f'"status": "{nouveau}"' + bloc[m2.end():]
        src = src[:m.start(1)] + nouveau_bloc + src[m.end(1):]
        changements.append((ident, ancien, nouveau))

    if changements:
        path.write_text(src, encoding="utf-8")

    print(f"inventory.js : {len(changements)} statut(s) modifié(s)")
    for ident, ancien, nouveau in changements:
        print(f"  - {ident} : {ancien} -> {nouveau}")

    if absents_du_fichier:
        print("Identifiants connus du catalogue mais absents d'inventory.js (ignorés) :")
        for i in absents_du_fichier: print("  -", i)

    if doublons:
        print("Identifiants présents plusieurs fois dans le tableau (dernière valeur conservée) :")
        for ident, occ in doublons.items():
            print(f"  - {ident} :", ", ".join(f"ligne {n} = {v!r}" for n, v in occ))

    if non_reconnus:
        print("Statuts texte non reconnus (ignorés, aucune modification) :")
        for n_ligne, ident, txt in non_reconnus:
            print(f"  - ligne {n_ligne} ({ident}) : {txt!r}")

    if inconnus:
        print("Identifiants du tableau absents du catalogue (ignorés) :")
        for n_ligne, i in inconnus:
            print(f"  - ligne {n_ligne} : {i}")

    manquants = sorted(connus - set(vus.keys()))
    if manquants:
        print("Modèles du catalogue sans ligne de statut dans le tableau (statut inchangé) :")
        for i in manquants: print("  -", i)

if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
